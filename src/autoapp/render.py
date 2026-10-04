"""Writes the CV and cover letter as Word documents."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .models import Profile

FONT = "Calibri"


def _doc() -> Document:
    d = Document()
    for s in d.sections:
        s.top_margin, s.bottom_margin = Cm(1.8), Cm(1.8)
        s.left_margin = s.right_margin = Cm(2.2)
    st = d.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(3)
    return d


def _heading(d: Document, text: str) -> None:
    p = d.add_paragraph()
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(10), Pt(3)
    r = p.add_run(text.upper())
    r.bold, r.font.size, r.font.color.rgb = True, Pt(10.5), RGBColor(0x1F, 0x3A, 0x5F)
    pb = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "4"), ("space", "1"), ("color", "1F3A5F")):
        bottom.set(qn(f"w:{k}"), v)
    pb.append(bottom)
    p._p.get_or_add_pPr().append(pb)


def _bullet(d: Document, text: str) -> None:
    p = d.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    p.add_run(text.rstrip("."))


def _contact_line(p: Profile) -> str:
    return "  |  ".join(x for x in [p.location, p.email, p.phone, *p.links] if x)


def _labels(lang: str) -> dict[str, str]:
    if lang == "de":
        return dict(profile="Profil", skills="Kenntnisse", exp="Berufserfahrung", edu="Ausbildung",
                    lang="Sprachen", interests="Interessen", present="heute", cv="Lebenslauf",
                    cover="Anschreiben")
    return dict(profile="Profile", skills="Skills", exp="Experience", edu="Education",
                lang="Languages", interests="Interests", present="present", cv="CV", cover="Cover letter")


def write_cv(path: Path, profile: Profile, content: dict, lang: str = "en") -> None:
    L = _labels(lang)
    d = _doc()
    p = d.add_paragraph()
    r = p.add_run(profile.name)
    r.bold, r.font.size = True, Pt(20)
    if profile.headline:
        d.add_paragraph(profile.headline)
    c = d.add_paragraph(_contact_line(profile))
    c.runs[0].font.size = Pt(9.5)

    if content.get("summary"):
        _heading(d, L["profile"])
        d.add_paragraph(content["summary"])
    if content.get("skills"):
        _heading(d, L["skills"])
        d.add_paragraph(", ".join(content["skills"]))

    by_idx = {e["index"]: e.get("bullets", []) for e in content.get("experience", []) if "index" in e}
    _heading(d, L["exp"])
    for i, e in enumerate(profile.experience):  # titles, employers and dates come from the profile
        end = L["present"] if e.end.lower() == "present" else e.end
        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.keep_with_next = True
        p.add_run(f"{e.title}, {e.company}").bold = True
        meta = " | ".join(x for x in [e.location, f"{e.start} to {end}"] if x)
        p.add_run("\n" + meta).italic = True
        for b in by_idx.get(i) or e.bullets:
            _bullet(d, b)

    by_edu = {e["index"]: e.get("details", []) for e in content.get("education", []) if "index" in e}
    _heading(d, L["edu"])
    for i, e in enumerate(profile.education):
        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.keep_with_next = True
        p.add_run(f"{e.degree}, {e.institution}").bold = True
        meta = " | ".join(x for x in [e.location, f"{e.start} to {e.end}" if e.start else ""] if x)
        if meta:
            p.add_run("\n" + meta).italic = True
        for dline in by_edu.get(i, e.details):
            _bullet(d, dline)

    for title, items in profile.extra_sections.items():
        _heading(d, title)
        for it in items:
            _bullet(d, it)
    if profile.languages:
        _heading(d, L["lang"])
        d.add_paragraph(", ".join(profile.languages))
    if profile.interests:
        _heading(d, L["interests"])
        d.add_paragraph(", ".join(profile.interests))
    d.save(path)


def letter_text(profile: Profile, job_company: str, content: dict, lang: str = "en") -> str:
    body = "\n\n".join(content.get("paragraphs", []))
    return (f"{content.get('subject', '')}\n\n{content.get('salutation', '')}\n\n{body}\n\n"
            f"{content.get('closing', '')}\n{profile.name}\n")


def write_letter_docx(path: Path, profile: Profile, company: str, content: dict, lang: str = "en") -> None:
    d = _doc()
    d.styles["Normal"].font.size = Pt(11)
    p = d.add_paragraph()
    p.add_run(profile.name).bold = True
    for line in [profile.location, profile.email, profile.phone]:
        if line:
            d.add_paragraph(line).paragraph_format.space_after = Pt(0)
    p = d.add_paragraph(date.today().strftime("%d.%m.%Y" if lang == "de" else "%d %B %Y"))
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d.add_paragraph(company).paragraph_format.space_before = Pt(6)
    s = d.add_paragraph()
    s.paragraph_format.space_before = Pt(14)
    s.add_run(content.get("subject", "")).bold = True
    d.add_paragraph(content.get("salutation", "")).paragraph_format.space_before = Pt(8)
    for para in content.get("paragraphs", []):
        q = d.add_paragraph(para)
        q.paragraph_format.space_after = Pt(8)
        q.paragraph_format.line_spacing = 1.15
    d.add_paragraph(content.get("closing", "")).paragraph_format.space_before = Pt(6)
    d.add_paragraph(profile.name)
    d.save(path)
